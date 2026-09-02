#!/usr/bin/env python3
"""
USCG Auxiliary Southeast District 7 HF Contingency Net Log
Python emulation of NetLog.xlsm + Netlog_Archives.xlsm

Email List worksheet → SQLite database (netlog.db)
All other reference data and formula logic embedded.

Usage:
    python netlog.py              # Interactive net log entry
    python netlog.py --new        # Start new net session
    python netlog.py --load FILE  # Load existing session JSON
    python netlog.py --archives   # View archive history
    python netlog.py --verify XLSX # Verify outputs against Excel file
"""

import json, math, os, sys, sqlite3
from datetime import datetime, date
from typing import Optional, Union
from dataclasses import dataclass, field, asdict

# ─────────────────────────────────────────────────────────────────────────────
# REFERENCE DATA  (FREQUENCIES, OP_CODE, TRANSCEIVERS, lookup lists)
# ─────────────────────────────────────────────────────────────────────────────

SECTORS = [
    "CHARLESTON", "JACKSONVILLE", "MIAMI",
    "KEY WEST", "SAN JUAN", "ST. PETERSBURG", "OTHER DISTRICTS"
]

NET_NAMES = ["SOUTHEAST DIGITAL NET", "SOUTHEAST VOICE NET", "SHARES SE NET"]
LOCATIONS = ["HOME", "REMOTE"]
PROPAGATIONS = ["EXCELLENT", "GOOD", "ACCEPTABLE", "POOR", "UNACCEPTABLE"]
NOISE_LEVELS = ["LOW:  S1-S3", "MODERATE:  S4-S6", "HIGH:  S5-S9 OR >"]
ANTENNAS = ["ANTENNA #1", "ANTENNA #2", "ANTENNA #3"]

TRANSCEIVERS = [
    "NONE","MANUAL","CODAN","FLEX RADIOS",
    "Icom 705","Icom 706","Icom 7100","Icom 7200","Icom 7300","Icom 7410",
    "Icom 7600","Icom 7610","Icom 9100","Icom IC-F8101",
    "Kenwood TS-590S","Kenwood TS-590SG","Kenwood TS-850S","Kenwood TS-890S",
    "Kenwood Comm.","Yaesu FT-100","Yaesu FT-450","Yaesu FT-600",
    "Yaesu FT-710","Yaesu FT-817","Yaesu FT-840","Yaesu FT-847",
    "Yaesu FT-857","Yaesu FT-891","Yaesu FT-897","Yaesu FT-920",
    "Yaesu FT-950","Yaesu FT-991/A","Yaesu FT-1000","Yaesu FT-2000",
    "Yaesu FTDX10","Yaesu FTdx101D/MP",
]

# Format: (code, khz_str, display_str)
FREQUENCIES = [
    ("NONE","NONE","NONE ASSIGNED"),
    ("A2A","2124.2","A2A:   2.1242 MHz"),("A2B","2479.0","A2B:   2.479 MHz"),
    ("A2C","2234.0","A2C:   2.234 MHz"),("A2D","2447.5","A2D:   2.4475 MHz"),
    ("A2E","2810.3","A2E:   2.8103 MHz"),("A3F","3154.5","A3F:   3.1545 MHz"),
    ("A3G","3203.0","A3G:   3.203 MHz"),("A3H","3390.5","A3H:   3.3905 MHz"),
    ("B4D","4084.8","B4D:   4.0848 MHz"),("B4E","4126.4","B4E:   4.1264 MHz"),
    ("B4A","4819.5","B4A:   4.8195 MHz"),("B4B","4847.0","B4B:   4.847 MHz"),
    ("B4C","4965.3","B4C:   4.9653 MHz"),("B5D","5253.5","B5D:   5.2535 MHz"),
    ("B5E","5322.5","B5E:   5.3225 MHz"),("B5F","5845.2","B5F:   5.8452 MHz"),
    ("B5G","5855.3","B5G:   5.8553 MHz"),("C6A","6821.8","C6A:   6.8218 MHz"),
    ("C6B","6972.8","C6B:   6.9728 MHz"),("C7C","7351.5","C7C:   7.3515 MHz"),
    ("C7D","7542.0","C7D:   7.542 MHz"),("D7A","7736.5","D7A:   7.7365 MHz"),
    ("D7B","7743.5","D7B:   7.7435 MHz"),("D8C","8002.3","D8C:   8.0023 MHz"),
    ("D8D","8035.3","D8D:   8.0353 MHz"),("E9A","9183.5","E9A:   9.1835 MHz"),
    ("E9B","9195.5","E9B:   9.1955 MHz"),("E9C","9338.5","E9C:   9.3385 MHz"),
    ("E10D","10508.3","E10D:   10.5083 MHz"),("F10A","10986.0","F10A:   10.986 MHz"),
    ("F11B","11030.5","F11B:   11.0305 MHz"),("F11C","11061.0","F11C:   11.061 MHz"),
    ("F11D","11115.0","F11D:   11.115 MHz"),("F11E","11130.3","F11E:   11.1303 MHz"),
    ("G14A","14459.5","G14A:   14.4595 MHz"),("G15B","15740.5","G15B:   15.7405 MHz"),
    ("G15C","15979.5","G15C:   15.9795 MHz"),("G16D","16128.5","G16D:   16.1285 MHz"),
    ("G16E","16210.0","G16E:   16.21 MHz"),
]
FREQ_DISPLAY = [f[2] for f in FREQUENCIES]

OP_CODES = [
    "NONE","VARA HF WINLINK","TELNET WINLINK","MT63-1KL","MT63-2KL",
    "OLIVIA OL 4-125","OLIVIA OL 4-250","OLIVIA OL 4-500","OLIVIA OL 4-1K","OLIVIA OL 4-2K",
    "OLIVIA OL 8-125","OLIVIA OL 8-250","OLIVIA OL 8-500","OLIVIA OL 8-1K","OLIVIA OL 8-2K",
]

# ─────────────────────────────────────────────────────────────────────────────
# EMAIL LIST DATABASE  (replaces hardcoded MEMBERS dict)
# Seed data extracted directly from NetLog.xlsm / Email List sheet
# Columns: callsign, operator_name, city, state, email, sector,
#          cell_phone, ham_callsign, shares_cs, winlink_cs, vara_hf_cs, notes, is_synthetic
# ─────────────────────────────────────────────────────────────────────────────

DB_PATH         = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'netlog.db')
ARCHIVE_DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'net_archives.db')
_db_conn: Optional[sqlite3.Connection]         = None
_archive_db_conn: Optional[sqlite3.Connection] = None

# Seed rows from Excel Email List (rows 2-74, excluding Reserved rows 75-90)
# Format: (callsign, operator_name, city, state, email, sector, cell_phone,
#          ham_callsign, shares_cs, winlink_cs, vara_hf_cs, notes, is_synthetic)
_SEED_MEMBERS = [
    # Synthetic entries (rows 2-3)
    ("NOTHING HEARD",  "NOTHING HEARD",  None, None, None, None, None, None, None, None, None, None, 1),
    ("STATION SECURED","STATION SECURED", None, None, None, None, None, None, None, None, None, None, 1),
    # Active/inactive members (rows 4-74)
    ("NF07BJ","Bigrow, John (NF07BJ)","Longs","SC","rowbig@gmail.com","CHARLESTON","631-786-2790","KD2CZD (G)",None,None,None,"No Winlink",0),
    ("NF07CE","Elliott, Bob (NF07CE)","Satsuma","FL","c.r.elliott@outlook.com","JACKSONVILLE","407-416-8410",None,None,None,None,"Status DTH",0),
    ("NF07CF","Miles, Albert (NF07CF)","Marietta","GA","amilesg@gmail.com","(CHARLESTON)",None,None,None,None,None,"Aux Retired",0),
    ("NF07CM","Miranda, Carmen (NF07CM)","Vero Beach","FL","cmirandak4crm","MIAMI","954-336-5190","K4CRM (G)",None,None,None,None,0),
    ("NF07CT","Towles, Chuck (NF07CT)","Pooler","GA","chuk1t@msn.com","CHARLESTON","202-550-0350","N4KKD (E)",None,None,None,"No Winlink",0),
    ("NF07DC","Creed, Dean (NF07DC)","Mt. Pleasant","SC","dean.creed@gmail.com","CHARLESTON","843-708-1541","K2DBC (G)",None,None,None,"No Winlink",0),
    ("NF07DE","Elliot, David (NF07DE)","Palm City","FL","david.elliot@cgauxnet.us","MIAMI","201-981-9908","K2DAE (A)",None,None,None,"No Winlink",0),
    ("NF07DR","Rockwell, David (NF07DR)","St. Pete Beach","FL","dave@daverockwell.com","ST. PETERSBURG","727-804-5821","W4PXE",None,None,None,"No Winlink",0),
    ("NF07DS-","Sullivan, Denis (NF07DS)","Melbourne","FL","dencats@att.net","(JACKSONVILLE)",None,None,None,None,None,"Aux Retired",0),
    ("NF07EQ","Quintela, Ed (NF07EQ)","Boca Raton","FL","equintel@bellsouth.net","MIAMI","305-582-5692","KN4NEI (E)","NF07EQ","NF07EQ","NF07EQ","Has Winlink & Vara Lic.",0),
    ("NF07GD","Durham, Gerald (NF07GD)","Estero","FL","gdurhamuscgaux91@gmai.com","ST. PETERSBURG","909-224-4517","KJ6PYL (T)",None,None,None,"No Winlink",0),
    ("NF07HM","Marschall, Harold (NF07HM)","Savannah","GA","hal.marschall@cgauxnet.us","CHARLESTON","516-840-2760","KO4AEM (T)",None,None,None,None,0),
    ("NF07JB","Blanchard, Joe (NF07JB)","Fernandina","FL","jlblan2@bellsouth.net","JACKSONVILLE","904-753-1007","WA4JLB (E)","NF07JB","NF07JB","NF07JB","Has Winlink & Vara Lic.",0),
    ("NF07JE","Burke, John (NF07JE)","Yulee","FL","mangoimg@bellsouth.net","JACKSONVILLE","904-583-4050","W0ENE (E)",None,None,None,"No Winlink",0),
    ("NF07JF","Froehler, John (NF07JF)","Jensen Beach","FL","johnfroehler@gmail.com","MIAMI","757-650-4275","K4LT (E)","NF07JF","NF07JF","NF07JF","Has Winlink & Vara Lic.",0),
    ("NF07JH","Herald, John (NF07JH)","St. Marys","GA","john.herald@tds.net","JACKSONVILLE","912-674-6479","KW4JMH (G)","NF07JH","NF07JH","NF07JH","Has Winlink & Vara Lic.",0),
    ("NF07JL","Laurent, John (NF07JL)","Barlow","FL","jfranklaurent@msn.com","MIAMI","863-255-4424",None,None,None,None,None,0),
    ("NF07JM","Mracna, Joe (NF07JM)","Seneca","SC","josephmracna@gmail.com","CHARLESTON","724-902-6662","AA3XE (E)","NF07JM","NF07JM","NF07JM","Has Winlink & Vara Lic.",0),
    ("NF07JS","Satterfield, Jack (NF07JS)","St. Pete Beach","FL","jack@satterfield.org","ST. PETERSBURG","727-403-5471",None,None,None,None,None,0),
    ("NF07JU","Nagy, Julius (NF07JU)","Melbourne","FL","juliusnagy1943@gmail.com","(JACKSONVILLE)",None,None,None,None,None,"Status DTH",0),
    ("NF07KB","Butler, Ken (NF07KB)","Hilton Head Island","SC","krbutler.mcgc@gmail.com","CHARLESTON","404-313-6272","N9QCV (G)","NF07KB","NF07KB",None,"No Winlink",0),
    ("NF07KH","FLOTILLA 5-8 (NF07KH)","Fort Pierce","FL","kwhontz1972@gmail.com","MIAMI","610-442-4970",None,None,None,None,"Handled by Karl Hontz",0),
    ("NF07LD","Lamy, Don (NF07LD)","Cooper City","FL","dlamy2@att.net","(MIAMI)",None,"KM2R (E)",None,None,None,"Resigned",0),
    ("NF07LP","Pernice, Lou (NF07LP)","Viera","FL","cgaux6799@gmail.com","(JACKSONVILLE)",None,"K4LRP (E)",None,None,None,"Aux Retired",0),
    ("NF07MG","Gaisford, Mark (NF07MG)","Stuart","FL","markgaisford@gmail.com","MIAMI","617-438-8686","W1MCG (G)","NF07MG","NF07MG","NF07MG",None,0),
    ("NF07OR","Thomas, Deborah (NF07OR)","Crystal River","FL","dthomasaux1501@gmail.com","ST. PETERSBURG","352-257-1829","K4FRH (E)",None,None,None,"No Winlink",0),
    ("NF07PP","Colee, Joe (NF07PP)","St. Augustine","FL","audioamp@aol.com","JACKSONVILLE","904-806-4877","KB4OLY (E)",None,"NF07PP",None,"Has Winlink",0),
    ("NF07RD","Damon, Rick (NF07RD)","Port Charlotte","FL","rdamon.uscgaux@gmail.com","ST. PETERSBURG","941-889-8659","K4KCR (G)",None,None,None,"No Winlink",0),
    ("NF07RR","Rideout, Richard (NF07RR)","Crossville","TN","richardrideout17@gmail.com","(OTHER DISTRICTS)",None,None,None,None,None,"Aux Retired",0),
    ("NF07SC","Schwartz, David (NF07SC)","Ft. Myers","FL","schwartzds@gmail.com","ST. PETERSBURG","239-898-0787","WX4S (E)",None,"NF07SC",None,"Has Winlink",0),
    ("NF07SJ","Shepard, Jim (NF07SJ)","Hilton Head","SC","shepardojimu@gmail.com","CHARLESTON","843-290-6233","K4RXH (G)","NF07SJ","NF07SJ",None,"Has Winlink",0),
    ("NF07SM","Magro, Sal (NF07SM)","Bradenton","FL","ddsl8@hotmail.com","ST. PETERSBURG","941-345-6419","N2UVG (G)",None,None,None,"No Winlink",0),
    ("NF07SR","FLOTILLA 5-9 (NF07SR)","Stuart","FL","flotilla59@gmail.com","MIAMI","757-650-4275",None,None,None,None,None,0),
    ("NF07TB","Baffa, Tom (NF07TB)","Amelia Island","FL","tbaffa@hotmail.com","JACKSONVILLE","913-787-3083","W0GEO (G)",None,None,None,"No Winlink",0),
    ("NF07WW","Walter, Bill (NF07WW)","Midway","GA","bllwltr@gmail.com","CHARLESTON","727-698-1382","N4DHP (G)","NF07WW","NF07WW","NF07WW","Has Winlink & Vara Lic.",0),
    ("NF07WY","Yawn, Bill (NF07WY)","Brunswick","GA","captain.bill.yawn@gmail.com","CHARLESTON","775-240-5144","WA4TFJ (A)",None,None,None,None,0),
    ("NM07EQ","Quintela, Ed (NM07EQ)","Boca Raton","FL","equintel@bellsouth.net","MIAMI","305-582-5692","KN4NEI (E)","NM07EQ","NM07EQ","NM07EQ","Has Winlink & Vara Lic.",0),
    ("NM07JA","Aleba, Joe (NM07JA)","Delray Beach","FL","jaleba@ntsci.com","MIAMI","954-214-5653",None,None,None,None,None,0),
    ("NM07MG","Gaisford, Mark (NM07MG)","Stuart","FL","markgaisford@gmail.com","MIAMI","617-438-8686","W1MCG (G)",None,None,None,"No Winlink",0),
    ("NM07MJ","Mracna, Joe (NM07MJ)","Seneca","SC","josephmracna@gmail.com","CHARLESTON","724-902-6662","AA3XE (E)","NF07JM","NF07JM","NF07JM","Has Winlink & Vara Lic.",0),
    ("NM07OT","Trott, Owen (NM07OT)","Charleston","SC","owen.trott.landman@gmail.com","CHARLESTON","936-645-5959","KF5BLK (G)",None,None,None,None,0),
    ("NM07TA","Alvord, Thomas (NM07TA)","Longboat Key","FL","alvord.uscgaux@gmail.com","ST. PETERSBURG","412-400-6166","AA3ZA (E)","NM07TA","NM07TA","NM07TA","Has Winlink",0),
    ("NT07AA","Gaisford, Mark (NT07AA)","Stuart","FL","markgaisford@gmail.com","MIAMI","617-438-8686","W1MCG (G)",None,None,None,"No Winlink",0),
    ("NT07AH","Hendrickson, Ames (NT07AH)","Fernandina","FL","ameshendrickson@gmail.com","JACKSONVILLE","541-480-8060","N7NPD (G)",None,None,None,None,0),
    ("NT07CA","Arenas, Carlos (NT07CA)","Lighthouse Pt","FL","webcapino@gmail.com","MIAMI","954-224-1655","W4BLQ (G)",None,None,None,"No Winlink",0),
    ("NT07DS","Smith, Dennis (NT07DS)","Orange Park","FL","des333@aol.com","MIAMI","904-813-2602",None,None,None,None,"No Winlink",0),
    ("NT07EM","Matos, Edwin (NT07EM)","Boca Raton","FL","semper.paratus.aux@gmail.com","MIAMI","786-316-7273","W2RLM (G)","NT07EM","NT07EM","NT07EM","Has Winlink & Vara Lic.",0),
    ("NT07EN","Nina, Emmanuel (NT07EN)","Deerfield Beach","FL","enina.cgaux@gmail.com","MIAMI","654-605-0204","W4ENS (E)","NT07EN","NT07EN","NT07EN","Has Winlink & Vara Lic.",0),
    ("NT07ES","Sora, Efrain (NT07ES)","Miami","FL","eisora1@yahoo.com","MIAMI","786-390-4663","KQ4DGC (T)",None,None,None,None,0),
    ("NT07GF","Fox, Gwen (NT07GF)","Bluffton","SC","drgwenfox@gmail.com","CHARLESTON","301-793-2551","KN4VWQ (E)",None,None,None,"No Winlink",0),
    ("NT07JJ","McKinley, Jamey (NT07JJ)","Bushnell","FL","uscgauxmckinley@gmail.com","JACKSONVILLE","321-327-1157","W4CGX (T)",None,None,None,None,0),
    ("NT07JK","Froehler, John (NT07JK)","Jensen Beach","FL","johnfroehler@gmail.com","MIAMI","757-650-4275","K4LT (E)",None,None,None,"No Winlink",0),
    ("NT07KH","Hontz, Karl (NT07KH)","Fort Pierce","FL","kwhontz1972@gmail.com","MIAMI","610-442-4970",None,"NT07KH","NT07KH",None,"Has Winlink",0),
    ("NT07MB","Bai, Mario (NT07MB)","Boca Raton","FL","cgauxbai@gmail.com","MIAMI","646-713-7272","N4BAI (G)","NT07MB","NT07MB","NT07MB","Has Winlink",0),
    ("NT07MK","Kline, Mike (NT07MK)","Boynton Beach","FL","klinemw@aol.com","MIAMI","913-620-4898",None,None,None,None,None,0),
    ("NT07MT","Alvord, Thomas (NT07MT- FL)","Longboat Key","FL","alvord.uscgaux@gmail.com","ST. PETERSBURG","412-400-6166","AA3ZA (E)",None,None,None,None,0),
    ("NT07PW","Ward, Phil (NT07PW)","Lehigh Acres","FL","kd4nyy@yahoo.com","ST. PETERSBURG","239-240-7741",None,None,None,None,None,0),
    ("NT07RW","Wannagot, Robert (NT07RW)","Port St. Lucie","FL","rwannagot@aol.com","MIAMI","203-260-1578","WA1GOT (G)",None,None,None,"No Winlink",0),
    ("NT07SD","Sabbagh, David (NT07SD)","Hillsboro Beach","FL","daytimer@comcast.net","MIAMI","847-284-6905","KN4TNA (G)","NT07SD","NT07SD","NT07SD","Has Winlink & Vara Lic.",0),
    ("NT07TH","Hammond, Ted (NT07TH)","West Palm Beach","FL","tnh@hammondfinancialservicesinc.com","MIAMI","636-448-7251",None,None,None,None,None,0),
    ("NF01FB","Brown, Fred (NF01FB)","Middle Grove","NY","w5bn12850@gmail.com","OTHER DISTRICTS","619-778-6210","W5BN (E)",None,None,None,None,0),
    ("NF01GR","Barbato, Greg (NF01GR)","Huntington","NY","gregbarbato@gmail.com","OTHER DISTRICTS","516-816-0847","AA2SX (E)",None,"NF07GR",None,"Has Winlink",0),
    ("NF01JC","Canavan, Jim (NF01JC)","Queensburry","NY","canavan@adelphia.net","OTHER DISTRICTS","518-796-5040",None,None,None,None,None,0),
    ("NF01RL","Light, Rich (NF01RL)","Bayville","NY","richlight1@gmail.com","OTHER DISTRICTS","516-528-1559","KD2RNN (E)",None,"NF07RL",None,"Has Winlink",0),
    ("NF013NG","Hopwood, William (NF013NG)","Elkins","NH","hopwoodwt@yahoo.com","OTHER DISTRICTS","603-526-6882","KB1QXJ (G)",None,None,None,None,0),
    ("NF81JK","Kenny, Jim (NF81JK)","Tallahassee","FL","jimc809@gmail.com","OTHER DISTRICTS","360-261-3193","KJ4KW (E)",None,None,None,None,0),
    ("NF82RN","Nagle, Roy (NF82RN)","Alabama","AL","ranagle@gmail.com","OTHER DISTRICTS","256-412-0447","KI4UX (E)",None,None,None,"No Winlink",0),
    ("NF85RH","Harden, Ray (NF85RH)","Clive","IA","uscgauxrh@outlook.com","OTHER DISTRICTS","515-208-0681","W0RAY (E)","NF85RH","NF85RH","NF85RH","Has Winlink & Vara Lic.",0),
    ("NF95CM","McNamara, Christopher (NF95CM)","Porter","IN","cmcnamara.uscgaux@gmail.com","OTHER DISTRICTS","219-628-3279","KC9VFR (G)","NF95CM","NF95CM","NF95CM","Has Winlink & Vara Lic.",0),
    ("NM81PF","Fenrich, Paul (NM81PF)","Cedar Park","TX","ka5fzu@austin.rr.com","OTHER DISTRICTS","512-923-7901","KA5FZU (T)",None,None,None,None,0),
]

_CREATE_TABLE = """
CREATE TABLE IF NOT EXISTS members (
    callsign        TEXT PRIMARY KEY NOT NULL,
    operator_name   TEXT,
    city            TEXT,
    state           TEXT,
    email           TEXT,
    sector          TEXT,
    cell_phone      TEXT,
    ham_callsign    TEXT,
    shares_cs       TEXT,
    winlink_cs      TEXT,
    vara_hf_cs      TEXT,
    notes           TEXT,
    is_synthetic    INTEGER NOT NULL DEFAULT 0,
    Latitude        REAL CHECK(Latitude  IS NULL OR (Latitude  >= 0 AND Latitude  <= 90)),
    Lat_Dir         TEXT CHECK(Lat_Dir   IS NULL OR Lat_Dir   IN ('N','S')),
    Longitude       REAL CHECK(Longitude IS NULL OR (Longitude >= 0 AND Longitude <= 180)),
    Long_Dir        TEXT CHECK(Long_Dir  IS NULL OR Long_Dir  IN ('E','W'))
);
CREATE INDEX IF NOT EXISTS idx_members_upper ON members (UPPER(callsign));
"""


def _init_db(db_path: str) -> sqlite3.Connection:
    conn = sqlite3.connect(db_path, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    for stmt in _CREATE_TABLE.strip().split(';'):
        if stmt.strip():
            conn.execute(stmt)
    conn.commit()

    # ── Migrate existing members table: add location columns if absent ────────
    for _sql in [
        "ALTER TABLE members ADD COLUMN Latitude  REAL CHECK(Latitude  IS NULL OR (Latitude  >= 0 AND Latitude  <= 90))",
        "ALTER TABLE members ADD COLUMN Lat_Dir   TEXT CHECK(Lat_Dir   IS NULL OR Lat_Dir   IN ('N','S'))",
        "ALTER TABLE members ADD COLUMN Longitude REAL CHECK(Longitude IS NULL OR (Longitude >= 0 AND Longitude <= 180))",
        "ALTER TABLE members ADD COLUMN Long_Dir  TEXT CHECK(Long_Dir  IS NULL OR Long_Dir  IN ('E','W'))",
    ]:
        try:
            conn.execute(_sql)
            conn.commit()
        except sqlite3.OperationalError:
            pass  # column already exists

    # ── members seed ─────────────────────────────────────────────────────────
    count = conn.execute("SELECT COUNT(*) FROM members").fetchone()[0]
    if count == 0:
        conn.executemany(
            """INSERT OR IGNORE INTO members
               (callsign, operator_name, city, state, email, sector,
                cell_phone, ham_callsign, shares_cs, winlink_cs, vara_hf_cs, notes, is_synthetic)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            _SEED_MEMBERS
        )
        conn.commit()

    # ── frequencies table ─────────────────────────────────────────────────────
    conn.execute("""
        CREATE TABLE IF NOT EXISTS frequencies (
            ID        INTEGER PRIMARY KEY,
            Frequency TEXT    NOT NULL
        )
    """)
    conn.commit()
    if conn.execute("SELECT COUNT(*) FROM frequencies").fetchone()[0] == 0:
        conn.executemany(
            "INSERT INTO frequencies (ID, Frequency) VALUES (?, ?)",
            [(i + 1, f[2]) for i, f in enumerate(FREQUENCIES)]
        )
        conn.commit()

    # ── sectors table ─────────────────────────────────────────────────────────
    conn.execute("""
        CREATE TABLE IF NOT EXISTS sectors (
            Sector TEXT PRIMARY KEY NOT NULL
        )
    """)
    conn.commit()
    if conn.execute("SELECT COUNT(*) FROM sectors").fetchone()[0] == 0:
        conn.executemany(
            "INSERT INTO sectors (Sector) VALUES (?)",
            [(s,) for s in SECTORS]
        )
        conn.commit()

    # ── transceivers table ────────────────────────────────────────────────────
    conn.execute("""
        CREATE TABLE IF NOT EXISTS transceivers (
            Transceiver TEXT PRIMARY KEY NOT NULL
        )
    """)
    conn.commit()
    if conn.execute("SELECT COUNT(*) FROM transceivers").fetchone()[0] == 0:
        conn.executemany(
            "INSERT INTO transceivers (Transceiver) VALUES (?)",
            [(t,) for t in TRANSCEIVERS]
        )
        conn.commit()

    # ── Net_Names table ───────────────────────────────────────────────────────
    conn.execute("""
        CREATE TABLE IF NOT EXISTS Net_Names (
            Nets TEXT PRIMARY KEY NOT NULL
        )
    """)
    conn.commit()
    if conn.execute("SELECT COUNT(*) FROM Net_Names").fetchone()[0] == 0:
        conn.executemany(
            "INSERT INTO Net_Names (Nets) VALUES (?)",
            [(n,) for n in NET_NAMES]
        )
        conn.commit()

    # ── Digital_Modes table ───────────────────────────────────────────────────
    conn.execute("""
        CREATE TABLE IF NOT EXISTS Digital_Modes (
            DigitalMode TEXT PRIMARY KEY NOT NULL
        )
    """)
    conn.commit()
    if conn.execute("SELECT COUNT(*) FROM Digital_Modes").fetchone()[0] == 0:
        conn.executemany(
            "INSERT INTO Digital_Modes (DigitalMode) VALUES (?)",
            [(m,) for m in OP_CODES]
        )
        conn.commit()

    return conn


def get_db() -> sqlite3.Connection:
    global _db_conn
    if _db_conn is None:
        _db_conn = _init_db(DB_PATH)
    return _db_conn


def get_freq_display() -> list:
    """Return the ordered frequency display strings from the frequencies table."""
    rows = get_db().execute(
        'SELECT Frequency FROM frequencies ORDER BY ID'
    ).fetchall()
    return [r[0] for r in rows]


def get_sectors() -> list:
    """Return the ordered sector names from the sectors table."""
    rows = get_db().execute(
        'SELECT Sector FROM sectors ORDER BY rowid'
    ).fetchall()
    return [r[0] for r in rows]


def get_transceivers() -> list:
    """Return the ordered transceiver names from the transceivers table."""
    rows = get_db().execute(
        'SELECT Transceiver FROM transceivers ORDER BY rowid'
    ).fetchall()
    return [r[0] for r in rows]


def get_net_names() -> list:
    """Return the ordered net names from the Net_Names table."""
    rows = get_db().execute(
        'SELECT Nets FROM Net_Names ORDER BY rowid'
    ).fetchall()
    return [r[0] for r in rows]


def get_digital_modes() -> list:
    """Return the ordered digital mode names from the Digital_Modes table."""
    rows = get_db().execute(
        'SELECT DigitalMode FROM Digital_Modes ORDER BY rowid'
    ).fetchall()
    return [r[0] for r in rows]


# ─────────────────────────────────────────────────────────────────────────────
# MATRIX LOADER  (loads BEARING/DISTANCE from the decrypted xlsx at startup)
# ─────────────────────────────────────────────────────────────────────────────

def load_matrices(xlsx_path: str) -> tuple:
    """Load BEARING and DISTANCE matrices. Returns (distance_matrix, bearing_matrix) as {(from,to): float}"""
    try:
        import openpyxl
    except ImportError:
        print("Warning: openpyxl not available. Distance/bearing lookups unavailable.")
        return {}, {}

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    result = {}
    for sheet_name in ['DISTANCE', 'BEARING']:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        headers = [str(v).strip() if v is not None else None for v in rows[0]]
        matrix = {}
        for row in rows[1:]:
            from_cs = str(row[0]).strip() if row[0] is not None else None
            if not from_cs:
                continue
            for col_idx, to_cs in enumerate(headers[1:], 1):
                if to_cs and col_idx < len(row) and row[col_idx] is not None:
                    try:
                        matrix[(from_cs, to_cs)] = float(row[col_idx])
                    except (ValueError, TypeError):
                        pass
        result[sheet_name] = matrix
    wb.close()
    return result.get('DISTANCE', {}), result.get('BEARING', {})


# ─────────────────────────────────────────────────────────────────────────────
# FORMULA ENGINE  (exact Python translations of every Excel formula)
# ─────────────────────────────────────────────────────────────────────────────

def _fuzzy_match(callsign: str, members: Optional[dict] = None) -> Optional[str]:
    """Implements Excel MATCH("*"&callsign, Email!A,) — exact then contains match.
    members=None uses DB; pass a dict to use a legacy dict (e.g. in verify harness)."""
    if not callsign:
        return None
    cs_upper = callsign.upper().strip()

    if members is not None:
        for k in members:
            if k.upper() == cs_upper:
                return k
        for k in members:
            if cs_upper in k.upper():
                return k
        return None

    conn = get_db()
    row = conn.execute(
        "SELECT callsign FROM members WHERE UPPER(callsign)=?", (cs_upper,)
    ).fetchone()
    if row:
        return row[0]
    row = conn.execute(
        "SELECT callsign FROM members WHERE INSTR(UPPER(callsign),?)>0 LIMIT 1", (cs_upper,)
    ).fetchone()
    return row[0] if row else None


def find_matching_callsigns(partial: str) -> list:
    """Return all members whose callsign contains `partial` (case-insensitive).
    Each item is a sqlite3.Row with callsign, operator_name, city, state, sector."""
    if not partial:
        return []
    cs_upper = partial.strip().upper()
    return get_db().execute(
        "SELECT callsign, operator_name, city, state, sector "
        "FROM members WHERE INSTR(UPPER(callsign),?)>0 ORDER BY callsign",
        (cs_upper,)
    ).fetchall()


def lookup_operator(callsign: str, members: Optional[dict] = None) -> str:
    """=IFERROR(INDEX(Email!B, MATCH("*"&D21, Email!A,)), "Callsign FROM Not Found")"""
    if members is not None:
        key = _fuzzy_match(callsign, members)
        return members[key].get('name', 'Callsign FROM Not Found') if key else 'Callsign FROM Not Found'
    key = _fuzzy_match(callsign)
    if not key:
        return 'Callsign FROM Not Found'
    row = get_db().execute("SELECT operator_name FROM members WHERE callsign=?", (key,)).fetchone()
    return (row['operator_name'] or 'Callsign FROM Not Found') if row else 'Callsign FROM Not Found'


def lookup_city(callsign: str, members: Optional[dict] = None) -> str:
    """=IFERROR(INDEX(Email!C, MATCH("*"&D21, Email!A,)), "CITY")"""
    if members is not None:
        key = _fuzzy_match(callsign, members)
        return (members[key].get('city') or 'CITY') if key else 'CITY'
    key = _fuzzy_match(callsign)
    if not key:
        return 'CITY'
    row = get_db().execute("SELECT city FROM members WHERE callsign=?", (key,)).fetchone()
    return (row['city'] or 'CITY') if row else 'CITY'


def lookup_state(callsign: str, members: Optional[dict] = None) -> str:
    """=IFERROR(INDEX(Email!D, MATCH("*"&D21, Email!A,)), "STATE")"""
    if members is not None:
        key = _fuzzy_match(callsign, members)
        return (members[key].get('state') or 'STATE') if key else 'STATE'
    key = _fuzzy_match(callsign)
    if not key:
        return 'STATE'
    row = get_db().execute("SELECT state FROM members WHERE callsign=?", (key,)).fetchone()
    return (row['state'] or 'STATE') if row else 'STATE'


def lookup_sector(callsign: str, members: Optional[dict] = None) -> str:
    """=IFERROR(INDEX(Email!F, MATCH("*"&D21, Email!A,)), "SECTOR")"""
    if members is not None:
        key = _fuzzy_match(callsign, members)
        return (members[key].get('sector') or 'SECTOR') if key else 'SECTOR'
    key = _fuzzy_match(callsign)
    if not key:
        return 'SECTOR'
    row = get_db().execute("SELECT sector FROM members WHERE callsign=?", (key,)).fetchone()
    return (row['sector'] or 'SECTOR') if row else 'SECTOR'


def lookup_distance(
    from_cs: str, to_cs: str,
    distance_matrix: dict,
    ncs_cs: str = "",
    show_my_db: bool = False
) -> Union[float, str]:
    """
    =IFNA(IF(AND(D21<>"",$D$23<>"",$A$25=TRUE),
          INDEX(DISTANCE!..., MATCH("*"&$D$23,...), MATCH("*"&D21,...)),
          IF(AND(ISBLANK(D21),ISBLANK(E21)),"ENTER FROM & TO",
            IF(ISBLANK(D21),"ENTER FROM",
              IF(ISBLANK(E21),"ENTER TO",
                INDEX(DISTANCE!..., MATCH("*"&D21,...), MATCH("*"&E21,...))
              )))), "Callsign TO Not Found")
    """
    def _lookup(f, t):
        v = distance_matrix.get((f, t))
        if v is not None:
            return v
        for (df, dt), dv in distance_matrix.items():
            if f.upper() in df.upper() and t.upper() in dt.upper():
                return dv
        return "Callsign TO Not Found"

    if show_my_db and from_cs and ncs_cs:
        return _lookup(ncs_cs, from_cs)
    if not from_cs and not to_cs:
        return "ENTER FROM & TO"
    if not from_cs:
        return "ENTER FROM"
    if not to_cs:
        return "ENTER TO"
    return _lookup(from_cs, to_cs)


def lookup_bearing(
    from_cs: str, to_cs: str,
    bearing_matrix: dict,
    ncs_cs: str = "",
    show_my_db: bool = False
) -> Union[float, str]:
    """Exact translation of BEARING formula — identical structure to distance."""
    def _lookup(f, t):
        v = bearing_matrix.get((f, t))
        if v is not None:
            return v
        for (bf, bt), bv in bearing_matrix.items():
            if f.upper() in bf.upper() and t.upper() in bt.upper():
                return bv
        return "Callsign TO Not Found"

    if show_my_db and from_cs and ncs_cs:
        return _lookup(ncs_cs, from_cs)
    if not from_cs and not to_cs:
        return "ENTER FROM & TO"
    if not from_cs:
        return "ENTER FROM"
    if not to_cs:
        return "ENTER TO"
    return _lookup(from_cs, to_cs)


def lookup_coords(callsign: str):
    """
    Return (signed_lat, signed_lon) from the members DB for *callsign*.

    Latitude:  N → positive,  S → negative.
    Longitude: E → positive,  W → negative.
    Returns (None, None) if the callsign is not found or has no coordinates.
    """
    if not callsign:
        return None, None
    key = _fuzzy_match(callsign)
    if not key:
        return None, None
    row = get_db().execute(
        'SELECT Latitude, Lat_Dir, Longitude, Long_Dir FROM members WHERE callsign=?',
        (key,)
    ).fetchone()
    if not row or row['Latitude'] is None or row['Longitude'] is None:
        return None, None
    lat = row['Latitude']  if row['Lat_Dir']  != 'S' else -row['Latitude']
    lon = row['Longitude'] if row['Long_Dir'] != 'W' else -row['Longitude']
    return lat, lon


def great_circle_distance(lat1: float, lon1: float,
                          lat2: float, lon2: float) -> float:
    """
    Haversine great-circle distance in statute miles.
    All arguments are signed decimal degrees (N/E positive, S/W negative).
    """
    R  = 3958.8                         # Earth radius in statute miles
    φ1 = math.radians(lat1)
    φ2 = math.radians(lat2)
    Δφ = math.radians(lat2 - lat1)
    Δλ = math.radians(lon2 - lon1)
    a  = (math.sin(Δφ / 2) ** 2
          + math.cos(φ1) * math.cos(φ2) * math.sin(Δλ / 2) ** 2)
    return R * 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))


def great_circle_bearing(lat1: float, lon1: float,
                         lat2: float, lon2: float) -> float:
    """
    Initial true bearing (0–360 °) on the great circle from point 1 to point 2.
    All arguments are signed decimal degrees.
    """
    φ1 = math.radians(lat1)
    φ2 = math.radians(lat2)
    Δλ = math.radians(lon2 - lon1)
    y  = math.sin(Δλ) * math.cos(φ2)
    x  = math.cos(φ1) * math.sin(φ2) - math.sin(φ1) * math.cos(φ2) * math.cos(Δλ)
    return (math.degrees(math.atan2(y, x)) + 360) % 360


def auto_timestamp(current_time: Optional[str], manual_mode: bool, callsign: str) -> Optional[str]:
    """=IF(AND($C$25=FALSE, D21<>""), IF(F21<>"",F21,NOW()),"")
    Preserves existing timestamp once set."""
    if manual_mode or not callsign:
        return None
    if current_time:
        return current_time
    return datetime.now().strftime('%H:%M:%S')


def rank_checkins(checkins: list) -> list:
    """=RANK(F21,$F$23:$F$91,1) — ascending rank by check-in time"""
    timestamped = [(i, c) for i, c in enumerate(checkins) if c.get('checkin_time')]
    timestamped.sort(key=lambda x: x[1]['checkin_time'])
    ranks = {orig_idx: rank for rank, (orig_idx, _) in enumerate(timestamped, 1)}
    for i, c in enumerate(checkins):
        c['checkin_order'] = ranks.get(i, '')
    return checkins


def count_by_sector(checkins: list) -> dict:
    """=COUNTIF($N$21:$N$91, "CHARLESTON") etc."""
    counts = {s: 0 for s in SECTORS}
    for c in checkins:
        sector = c.get('sector', '')
        if sector in counts:
            counts[sector] += 1
    counts['TOTAL'] = sum(counts[s] for s in SECTORS)
    return counts


def format_freq_display(code: str) -> str:
    """=A2 & ":   " & B2/1000 & " MHz"  — FREQUENCIES sheet formula"""
    for f in FREQUENCIES:
        if f[0] == code:
            return f[2]
    return code


# ─────────────────────────────────────────────────────────────────────────────
# DATA MODEL
# ─────────────────────────────────────────────────────────────────────────────

@dataclass
class CheckIn:
    """One station check-in (maps to the FROM+TO row pair in Excel)"""
    from_callsign: str = ""
    to_callsign: str = ""
    checkin_time: str = ""
    manual_time: bool = False
    signal_L: bool = False
    signal_G: bool = False
    signal_W: bool = False
    signal_VW: bool = False
    signal_F: bool = False
    signal_C: bool = False
    signal_R: bool = False
    signal_UR: bool = False
    signal_D: bool = False
    signal_WI: bool = False
    signal_I: bool = False
    signal_NH: bool = False
    has_traffic: bool = False
    traffic_destination: str = ""
    traffic_notes: str = ""
    propagation: str = ""        # per-check-in propagation (alt-freq conditions vary per station)
    # Formula-computed fields
    from_operator: str = ""
    to_operator: str = ""
    city: str = ""
    state: str = ""
    sector: str = ""
    distance_miles: Union[float, str] = ""
    bearing_degrees: Union[float, str] = ""
    checkin_order: Union[int, str] = ""
    dest_operator: str = ""

    def compute_auto_fields(self, session: 'NetSession', dist_matrix: dict, bear_matrix: dict):
        self.from_operator = lookup_operator(self.from_callsign)
        self.to_operator   = lookup_operator(self.to_callsign)
        self.city          = lookup_city(self.from_callsign)
        self.state         = lookup_state(self.from_callsign)
        self.sector        = lookup_sector(self.from_callsign)
        self.dest_operator = lookup_operator(self.traffic_destination) if self.traffic_destination else ""

        # ── Distance / bearing ─────────────────────────────────────────────────
        # Prefer great-circle calculation from DB coordinates; fall back to the
        # Excel-loaded matrices when either station has no coordinates on file.
        gc_done = False
        if self.from_callsign:
            # show_my_db swaps the pair so NCS→station distance is shown
            if session.show_my_db and session.ncs_callsign:
                cs_a, cs_b = session.ncs_callsign, self.from_callsign
            else:
                cs_a, cs_b = self.from_callsign, self.to_callsign

            if cs_a and cs_b:
                lat1, lon1 = lookup_coords(cs_a)
                lat2, lon2 = lookup_coords(cs_b)
                if lat1 is not None and lat2 is not None:
                    self.distance_miles  = great_circle_distance(lat1, lon1, lat2, lon2)
                    self.bearing_degrees = great_circle_bearing(lat1, lon1, lat2, lon2)
                    gc_done = True

        if not gc_done:
            self.distance_miles  = lookup_distance(
                self.from_callsign, self.to_callsign, dist_matrix,
                ncs_cs=session.ncs_callsign, show_my_db=session.show_my_db
            )
            self.bearing_degrees = lookup_bearing(
                self.from_callsign, self.to_callsign, bear_matrix,
                ncs_cs=session.ncs_callsign, show_my_db=session.show_my_db
            )

    def signal_summary(self) -> str:
        codes = [c for c in ['L','G','W','VW','F','C','R','UR','D','WI','I'] if getattr(self, f'signal_{c}', False)]
        return "/".join(codes)


@dataclass
class NetSession:
    """One complete net session (header block + check-in table)"""
    freq_type: str = "PRIMARY"
    net_name: str = ""
    net_date: str = ""
    start_time: str = ""
    transceiver: str = ""
    antenna: str = "ANTENNA #1"
    location: str = "HOME"
    primary_freq: str = "NONE ASSIGNED"
    alt_freq_1: str = "NONE ASSIGNED"
    alt_freq_2: str = "NONE ASSIGNED"
    digital_mode: str = "NONE"
    secured_time: str = ""
    propagation: str = ""
    noise_level: str = ""
    antenna_1_desc: str = "ICOM AH-710 FOLDED DIPOLE"
    antenna_2_desc: str = "BUSHCOMM SWE-100 BROADBAND SINGLE WIRE END FED"
    ncs_callsign: str = ""
    station_callsign: str = ""
    ancs_callsign: str = ""
    show_my_db: bool = False
    manual_time: bool = False
    checkins: list = field(default_factory=list)           # PRIMARY log
    checkins_alt1: list = field(default_factory=list)      # 1ST ALT log
    checkins_alt2: list = field(default_factory=list)      # 2ND ALT log

    # ── Per-frequency helpers ─────────────────────────────────────────────────
    def get_active_checkins(self) -> list:
        """Return the check-in list for the currently active frequency."""
        if self.freq_type == '1ST ALT':
            return self.checkins_alt1
        if self.freq_type == '2ND ALT':
            return self.checkins_alt2
        return self.checkins

    def all_checkins(self) -> list:
        """All check-ins across all three frequency logs (for archive counts etc.)."""
        return self.checkins + self.checkins_alt1 + self.checkins_alt2

    @property
    def sector_counts(self) -> dict:
        return count_by_sector(self.get_active_checkins())

    def add_checkin(self, dist_matrix: dict, bear_matrix: dict, **kwargs) -> 'CheckIn':
        ci = CheckIn(**{k: v for k, v in kwargs.items() if k in CheckIn.__dataclass_fields__})
        if not ci.manual_time and ci.from_callsign and not ci.checkin_time:
            ci.checkin_time = datetime.now().strftime('%H:%M:%S')
        ci.compute_auto_fields(self, dist_matrix, bear_matrix)
        lst = self.get_active_checkins()
        lst.append(ci.__dict__)
        rank_checkins(lst)
        return ci

    def recompute_all(self, dist_matrix: dict, bear_matrix: dict):
        for lst in [self.checkins, self.checkins_alt1, self.checkins_alt2]:
            for ci_dict in lst:
                ci = CheckIn(**{k: ci_dict.get(k, v.default if hasattr(v, 'default') else '')
                               for k, v in CheckIn.__dataclass_fields__.items()})
                ci.compute_auto_fields(self, dist_matrix, bear_matrix)
                ci_dict.update(ci.__dict__)
            rank_checkins(lst)

    def to_dict(self) -> dict:
        return {
            'freq_type': self.freq_type, 'net_name': self.net_name,
            'net_date': self.net_date, 'start_time': self.start_time,
            'transceiver': self.transceiver, 'antenna': self.antenna,
            'location': self.location, 'primary_freq': self.primary_freq,
            'alt_freq_1': self.alt_freq_1, 'alt_freq_2': self.alt_freq_2,
            'digital_mode': self.digital_mode, 'secured_time': self.secured_time,
            'propagation': self.propagation, 'noise_level': self.noise_level,
            'antenna_1_desc': self.antenna_1_desc, 'antenna_2_desc': self.antenna_2_desc,
            'ncs_callsign': self.ncs_callsign,
            'station_callsign': self.station_callsign, 'ancs_callsign': self.ancs_callsign,
            'show_my_db': self.show_my_db,
            'manual_time': self.manual_time,
            'checkins': self.checkins,
            'checkins_alt1': self.checkins_alt1,
            'checkins_alt2': self.checkins_alt2,
        }

    def save(self, path: str):
        with open(path, 'w') as f:
            json.dump(self.to_dict(), f, indent=2, default=str)

    @classmethod
    def load(cls, path: str) -> 'NetSession':
        with open(path) as f:
            data = json.load(f)
        s = cls()
        _checkin_keys = {'checkins', 'checkins_alt1', 'checkins_alt2'}
        for k, v in data.items():
            if k not in _checkin_keys:
                setattr(s, k, v)
        s.checkins      = data.get('checkins',      [])
        s.checkins_alt1 = data.get('checkins_alt1', [])
        s.checkins_alt2 = data.get('checkins_alt2', [])
        return s


# ─────────────────────────────────────────────────────────────────────────────
# ARCHIVE MANAGER
# ─────────────────────────────────────────────────────────────────────────────

ARCHIVE_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'netlog_archives.json')


# ── Archive DB helpers ────────────────────────────────────────────────────────

def _insert_checkin_row(conn: sqlite3.Connection,
                        session_id: int, freq_label: str, ci: dict):
    """Insert one check-in dict into the archive checkins table."""
    conn.execute("""
        INSERT INTO checkins (
            session_id, freq_type,
            from_callsign, to_callsign, checkin_time, manual_time,
            noise_level, antenna,
            signal_L, signal_G, signal_W, signal_VW, signal_F,
            signal_C, signal_R, signal_UR, signal_D, signal_WI, signal_I, signal_NH,
            has_traffic, traffic_destination, traffic_notes, propagation,
            from_operator, to_operator, city, state, sector,
            distance_miles, bearing_degrees, checkin_order, dest_operator
        ) VALUES (
            ?,?,
            ?,?,?,?,
            ?,?,
            ?,?,?,?,?,
            ?,?,?,?,?,?,?,
            ?,?,?,?,
            ?,?,?,?,?,
            ?,?,?,?
        )
    """, (
        session_id, freq_label,
        ci.get('from_callsign', ''), ci.get('to_callsign', ''),
        ci.get('checkin_time', ''), int(bool(ci.get('manual_time', False))),
        ci.get('noise_level', ''), ci.get('antenna', ''),
        int(bool(ci.get('signal_L',  False))),
        int(bool(ci.get('signal_G',  False))),
        int(bool(ci.get('signal_W',  False))),
        int(bool(ci.get('signal_VW', False))),
        int(bool(ci.get('signal_F',  False))),
        int(bool(ci.get('signal_C',  False))),
        int(bool(ci.get('signal_R',  False))),
        int(bool(ci.get('signal_UR', False))),
        int(bool(ci.get('signal_D',  False))),
        int(bool(ci.get('signal_WI', False))),
        int(bool(ci.get('signal_I',  False))),
        int(bool(ci.get('signal_NH', False))),
        int(bool(ci.get('has_traffic', False))),
        ci.get('traffic_destination', ''), ci.get('traffic_notes', ''),
        ci.get('propagation', ''),
        ci.get('from_operator', ''), ci.get('to_operator', ''),
        ci.get('city', ''), ci.get('state', ''), ci.get('sector', ''),
        ci.get('distance_miles', None), ci.get('bearing_degrees', None),
        ci.get('checkin_order', None), ci.get('dest_operator', ''),
    ))


def _init_archive_db(path: str) -> sqlite3.Connection:
    conn = sqlite3.connect(path, check_same_thread=False)
    conn.execute('PRAGMA journal_mode=WAL')
    conn.row_factory = sqlite3.Row

    conn.execute("""
        CREATE TABLE IF NOT EXISTS sessions (
            session_id    INTEGER PRIMARY KEY AUTOINCREMENT,
            freq_type     TEXT,
            net_name      TEXT,
            net_date      TEXT,
            start_time    TEXT,
            transceiver   TEXT,
            antenna       TEXT,
            location      TEXT,
            primary_freq  TEXT,
            alt_freq_1    TEXT,
            alt_freq_2    TEXT,
            digital_mode  TEXT,
            secured_time  TEXT,
            propagation   TEXT,
            noise_level   TEXT,
            antenna_1_desc TEXT,
            antenna_2_desc TEXT,
            ncs_callsign     TEXT,
            station_callsign TEXT,
            ancs_callsign    TEXT,
            show_my_db    INTEGER DEFAULT 0,
            manual_time   INTEGER DEFAULT 0,
            archived_at   TEXT
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS checkins (
            checkin_id          INTEGER PRIMARY KEY AUTOINCREMENT,
            session_id          INTEGER NOT NULL,
            freq_type           TEXT,
            from_callsign       TEXT,
            to_callsign         TEXT,
            checkin_time        TEXT,
            manual_time         INTEGER DEFAULT 0,
            noise_level         TEXT,
            antenna             TEXT,
            signal_L            INTEGER DEFAULT 0,
            signal_G            INTEGER DEFAULT 0,
            signal_W            INTEGER DEFAULT 0,
            signal_VW           INTEGER DEFAULT 0,
            signal_F            INTEGER DEFAULT 0,
            signal_C            INTEGER DEFAULT 0,
            signal_R            INTEGER DEFAULT 0,
            signal_UR           INTEGER DEFAULT 0,
            signal_D            INTEGER DEFAULT 0,
            signal_WI           INTEGER DEFAULT 0,
            signal_I            INTEGER DEFAULT 0,
            signal_NH           INTEGER DEFAULT 0,
            has_traffic         INTEGER DEFAULT 0,
            traffic_destination TEXT,
            traffic_notes       TEXT,
            propagation         TEXT,
            from_operator       TEXT,
            to_operator         TEXT,
            city                TEXT,
            state               TEXT,
            sector              TEXT,
            distance_miles      REAL,
            bearing_degrees     REAL,
            checkin_order       INTEGER,
            dest_operator       TEXT
        )
    """)
    conn.commit()

    # ── Migrate existing archive checkins table: add signal_NH if absent ─────
    try:
        conn.execute("ALTER TABLE checkins ADD COLUMN signal_NH INTEGER DEFAULT 0")
        conn.commit()
    except sqlite3.OperationalError:
        pass  # column already exists

    # ── Migrate existing sessions table: add station/ancs callsigns if absent ─
    for col in ('station_callsign TEXT', 'ancs_callsign TEXT'):
        try:
            conn.execute(f"ALTER TABLE sessions ADD COLUMN {col}")
            conn.commit()
        except sqlite3.OperationalError:
            pass  # column already exists

    # ── Migrate existing JSON archives on first run ───────────────────────────
    if conn.execute('SELECT COUNT(*) FROM sessions').fetchone()[0] == 0:
        if os.path.exists(ARCHIVE_FILE):
            try:
                with open(ARCHIVE_FILE) as fj:
                    old = json.load(fj)
                for entry in old:
                    cur = conn.execute("""
                        INSERT INTO sessions (
                            freq_type, net_name, net_date, start_time,
                            transceiver, antenna, location,
                            primary_freq, alt_freq_1, alt_freq_2,
                            digital_mode, secured_time, propagation, noise_level,
                            antenna_1_desc, antenna_2_desc,
                            ncs_callsign, station_callsign, ancs_callsign,
                            show_my_db, manual_time, archived_at
                        ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                    """, (
                        entry.get('freq_type', ''), entry.get('net_name', ''),
                        entry.get('net_date', ''), entry.get('start_time', ''),
                        entry.get('transceiver', ''), entry.get('antenna', ''),
                        entry.get('location', ''), entry.get('primary_freq', ''),
                        entry.get('alt_freq_1', ''), entry.get('alt_freq_2', ''),
                        entry.get('digital_mode', ''), entry.get('secured_time', ''),
                        entry.get('propagation', ''), entry.get('noise_level', ''),
                        entry.get('antenna_1_desc', ''), entry.get('antenna_2_desc', ''),
                        entry.get('ncs_callsign', ''),
                        entry.get('station_callsign', ''), entry.get('ancs_callsign', ''),
                        int(bool(entry.get('show_my_db', False))),
                        int(bool(entry.get('manual_time', False))),
                        entry.get('archived_at', ''),
                    ))
                    sid = cur.lastrowid
                    for label, key in [
                        ('PRIMARY', 'checkins'),
                        ('1ST ALT', 'checkins_alt1'),
                        ('2ND ALT', 'checkins_alt2'),
                    ]:
                        for ci in entry.get(key, []):
                            _insert_checkin_row(conn, sid, label, ci)
                conn.commit()
                print(f"Migrated {len(old)} JSON archive(s) to net_archives.db")
            except Exception as exc:
                print(f"Warning: JSON archive migration failed: {exc}")

    return conn


def get_archive_db() -> sqlite3.Connection:
    global _archive_db_conn
    if _archive_db_conn is None:
        _archive_db_conn = _init_archive_db(ARCHIVE_DB_PATH)
    return _archive_db_conn


def save_to_archive(session: NetSession):
    conn        = get_archive_db()
    archived_at = datetime.now().isoformat()

    cur = conn.execute("""
        INSERT INTO sessions (
            freq_type, net_name, net_date, start_time,
            transceiver, antenna, location,
            primary_freq, alt_freq_1, alt_freq_2,
            digital_mode, secured_time, propagation, noise_level,
            antenna_1_desc, antenna_2_desc,
            ncs_callsign, station_callsign, ancs_callsign,
            show_my_db, manual_time, archived_at
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
    """, (
        session.freq_type, session.net_name, session.net_date, session.start_time,
        session.transceiver, session.antenna, session.location,
        session.primary_freq, session.alt_freq_1, session.alt_freq_2,
        session.digital_mode, session.secured_time, session.propagation, session.noise_level,
        session.antenna_1_desc, session.antenna_2_desc,
        session.ncs_callsign, session.station_callsign, session.ancs_callsign,
        int(session.show_my_db), int(session.manual_time), archived_at,
    ))
    session_id = cur.lastrowid

    for freq_label, ci_list in [
        ('PRIMARY', session.checkins),
        ('1ST ALT', session.checkins_alt1),
        ('2ND ALT', session.checkins_alt2),
    ]:
        for ci in ci_list:
            if ci.get('pending'):   # skip NCS auto-entry if never confirmed
                continue
            _insert_checkin_row(conn, session_id, freq_label, ci)
    conn.commit()

    total = len(session.checkins) + len(session.checkins_alt1) + len(session.checkins_alt2)
    print(f"Session archived to net_archives.db  "
          f"(session_id={session_id}, {total} check-in(s))")


def load_archives() -> list:
    """Load all archived sessions from net_archives.db.
    Returns a list of dicts in the same format as the old JSON archive,
    so ArchivesTab and the CLI print_archives() work without change.
    """
    conn     = get_archive_db()
    sessions = conn.execute(
        'SELECT * FROM sessions ORDER BY session_id'
    ).fetchall()
    result   = []
    for s in sessions:
        entry = dict(s)
        entry['show_my_db']  = bool(entry.get('show_my_db',  0))
        entry['manual_time'] = bool(entry.get('manual_time', 0))

        all_cis = conn.execute(
            'SELECT * FROM checkins WHERE session_id=? ORDER BY checkin_order',
            (s['session_id'],)
        ).fetchall()

        checkins, checkins_alt1, checkins_alt2 = [], [], []
        for ci in all_cis:
            ci_dict = dict(ci)
            ft = ci_dict.pop('freq_type', 'PRIMARY')
            ci_dict.pop('checkin_id',  None)
            ci_dict.pop('session_id',  None)
            # Restore booleans
            for k in list(ci_dict):
                if k.startswith('signal_') or k in ('manual_time', 'has_traffic'):
                    ci_dict[k] = bool(ci_dict[k])
            if ft == '1ST ALT':
                checkins_alt1.append(ci_dict)
            elif ft == '2ND ALT':
                checkins_alt2.append(ci_dict)
            else:
                checkins.append(ci_dict)

        entry['checkins']      = checkins
        entry['checkins_alt1'] = checkins_alt1
        entry['checkins_alt2'] = checkins_alt2
        result.append(entry)
    return result


# ─────────────────────────────────────────────────────────────────────────────
# MEMBER MANAGEMENT (CRUD on netlog.db)
# ─────────────────────────────────────────────────────────────────────────────

def _cmd_browse_members():
    conn = get_db()
    rows = conn.execute(
        """SELECT callsign, operator_name, city, state, sector, winlink_cs, vara_hf_cs
           FROM members WHERE is_synthetic=0
           ORDER BY UPPER(callsign)"""
    ).fetchall()
    print(f"\n{'CALLSIGN':10} {'NAME':38} {'CITY':16} {'ST':3} {'SECTOR':16} {'WL':6} {'VARA':4}")
    print("-" * 100)
    for r in rows:
        wl = 'YES' if r['winlink_cs'] else ''
        vara = 'YES' if r['vara_hf_cs'] else ''
        print(f"{(r['callsign'] or ''):10} {(r['operator_name'] or ''):38} "
              f"{(r['city'] or ''):16} {(r['state'] or ''):3} {(r['sector'] or ''):16} {wl:6} {vara:4}")
    print(f"\n{len(rows)} members in database.")


def _cmd_new_member():
    print("\n--- NEW MEMBER ---")
    callsign = input("  Callsign (e.g. NF07XX): ").strip().upper()
    if not callsign:
        print("  Cancelled.")
        return
    existing = get_db().execute(
        "SELECT callsign FROM members WHERE UPPER(callsign)=?", (callsign.upper(),)
    ).fetchone()
    if existing:
        print(f"  {callsign} already exists — use [X] to edit.")
        return
    name   = input("  Operator name (Lastname, First (CS)): ").strip()
    city   = input("  City: ").strip()
    state  = input("  State (2-letter): ").strip().upper()[:2]
    email  = input("  Email: ").strip()
    sector = _prompt("Sector", choices=SECTORS + [f"({s})" for s in SECTORS])
    phone  = input("  Cell phone: ").strip()
    ham_cs = input("  Ham callsign (e.g. WA4JLB (E)): ").strip()
    wl_cs  = input("  Winlink callsign (blank=none): ").strip() or None
    vara   = input("  VARA HF licensed? (y/N): ").strip().lower() == 'y'
    shares = input("  SHARES member callsign (blank=none): ").strip() or None
    notes  = input("  Notes: ").strip()
    get_db().execute(
        """INSERT INTO members
           (callsign, operator_name, city, state, email, sector,
            cell_phone, ham_callsign, shares_cs, winlink_cs, vara_hf_cs, notes, is_synthetic)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,0)""",
        (callsign, name, city, state, email, sector, phone, ham_cs,
         shares, wl_cs, callsign if vara else None, notes)
    )
    get_db().commit()
    print(f"  {callsign} — {name} added.")


def _cmd_edit_member():
    print("\n--- EDIT MEMBER ---")
    callsign = input("  Callsign to edit: ").strip().upper()
    conn = get_db()
    row = conn.execute("SELECT * FROM members WHERE UPPER(callsign)=?", (callsign.upper(),)).fetchone()
    if not row:
        print(f"  {callsign} not found.")
        return
    if row['is_synthetic']:
        print(f"  {callsign} is a system entry — cannot edit.")
        return
    print(f"  Editing {row['callsign']} — {row['operator_name']}")
    print("  (Press ENTER to keep current value)\n")

    def ask(label, current):
        v = input(f"  {label} [{current or ''}]: ").strip()
        return v if v else current

    name   = ask("Operator name",    row['operator_name'])
    city   = ask("City",             row['city'])
    state  = ask("State",            row['state'])
    email  = ask("Email",            row['email'])
    sector = ask("Sector",           row['sector'])
    phone  = ask("Cell phone",       row['cell_phone'])
    ham_cs = ask("Ham callsign",     row['ham_callsign'])
    wl_cs  = ask("Winlink callsign", row['winlink_cs'])
    vara   = ask("VARA HF (cs/blank)", row['vara_hf_cs'])
    shares = ask("SHARES cs/blank",  row['shares_cs'])
    notes  = ask("Notes",            row['notes'])

    conn.execute(
        """UPDATE members SET
           operator_name=?, city=?, state=?, email=?, sector=?,
           cell_phone=?, ham_callsign=?, winlink_cs=?, vara_hf_cs=?, shares_cs=?, notes=?
           WHERE callsign=?""",
        (name, city, state, email, sector, phone, ham_cs,
         wl_cs or None, vara or None, shares or None, notes, row['callsign'])
    )
    conn.commit()
    print(f"  {row['callsign']} updated.")


def _cmd_delete_member():
    print("\n--- DELETE MEMBER ---")
    callsign = input("  Callsign to delete: ").strip().upper()
    conn = get_db()
    row = conn.execute(
        "SELECT callsign, operator_name, is_synthetic FROM members WHERE UPPER(callsign)=?",
        (callsign.upper(),)
    ).fetchone()
    if not row:
        print(f"  {callsign} not found.")
        return
    if row['is_synthetic']:
        print(f"  {callsign} is a system entry — cannot delete.")
        return
    confirm = input(
        f"  Delete {row['callsign']} ({row['operator_name']})?\n"
        f"  Re-type callsign to confirm: "
    ).strip().upper()
    if confirm != row['callsign'].upper():
        print("  Delete cancelled.")
        return
    conn.execute("DELETE FROM members WHERE callsign=?", (row['callsign'],))
    conn.commit()
    print(f"  {row['callsign']} deleted.")


# ─────────────────────────────────────────────────────────────────────────────
# VERIFICATION MODULE
# ─────────────────────────────────────────────────────────────────────────────

def verify_against_excel(xlsx_path: str):
    """Loads decrypted xlsx and verifies Python outputs match Excel cell values."""
    try:
        import openpyxl
    except ImportError:
        print("openpyxl required for verification")
        return

    dist_matrix, bear_matrix = load_matrices(xlsx_path)
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    print("\n=== VERIFICATION REPORT ===\n")
    tests_run = tests_passed = tests_failed = 0

    def check(name, got, expected):
        nonlocal tests_run, tests_passed, tests_failed
        tests_run += 1
        if isinstance(expected, float) and isinstance(got, (int, float)):
            ok = abs(float(got) - expected) < 0.01
        else:
            ok = str(got).strip() == str(expected).strip()
        if not ok:
            tests_failed += 1
            print(f"  [FAIL] {name}: got {got!r}, expected {expected!r}")
        else:
            tests_passed += 1
        return ok

    # Build legacy dict from Excel Email List for lookup testing
    ws_email = wb['Email List']
    excel_members = {}
    for row in ws_email.iter_rows(min_row=2, max_row=90, values_only=True):
        cs = str(row[0]).strip() if row[0] else None
        if not cs or cs.startswith('Reserved'):
            continue
        excel_members[cs] = {
            'name':       str(row[1]).strip() if row[1] else None,
            'city':       str(row[2]).strip() if row[2] else None,
            'state':      str(row[3]).strip() if row[3] else None,
            'email':      str(row[4]).strip() if row[4] else None,
            'sector':     str(row[5]).strip() if row[5] else None,
            'winlink_cs': str(row[9]).strip() if row[9] else None,
            'notes':      str(row[11]).strip() if row[11] else None,
        }

    # 1. Member lookup formulas (using Excel dict as ground truth)
    print("1. Member Lookup Formulas (Email List INDEX/MATCH):")
    test_rows = [r for r in list(ws_email.iter_rows(min_row=4, max_row=90, values_only=True))
                 if r[0] and not str(r[0]).startswith('Reserved')][:12]
    for row in test_rows:
        cs = str(row[0]).strip()
        if row[1]:
            check(f"lookup_operator({cs!r})", lookup_operator(cs, excel_members), str(row[1]).strip())
        if row[2]:
            check(f"lookup_city({cs!r})",     lookup_city(cs, excel_members),     str(row[2]).strip())
        if row[5]:
            check(f"lookup_sector({cs!r})",   lookup_sector(cs, excel_members),   str(row[5]).strip())

    # 2. Distance matrix
    print("\n2. Distance Matrix Lookups:")
    ws_dist = wb['DISTANCE']
    dist_rows = list(ws_dist.iter_rows(values_only=True))
    dist_headers = [str(v).strip() if v is not None else None for v in dist_rows[0]]
    for row in dist_rows[1:8]:
        from_cs = str(row[0]).strip() if row[0] else None
        if not from_cs or from_cs in ('NOTHING HEARD', 'STATION SECURED'):
            continue
        for col_idx, to_cs in enumerate(dist_headers[1:6], 1):
            if to_cs and to_cs not in ('NOTHING HEARD', 'STATION SECURED') and col_idx < len(row) and row[col_idx] is not None:
                check(f"distance({from_cs},{to_cs})", lookup_distance(from_cs, to_cs, dist_matrix), float(row[col_idx]))

    # 3. Bearing matrix
    print("\n3. Bearing Matrix Lookups:")
    ws_bear = wb['BEARING']
    bear_rows = list(ws_bear.iter_rows(values_only=True))
    bear_headers = [str(v).strip() if v is not None else None for v in bear_rows[0]]
    for row in bear_rows[1:5]:
        from_cs = str(row[0]).strip() if row[0] else None
        if not from_cs or from_cs in ('NOTHING HEARD', 'STATION SECURED'):
            continue
        for col_idx, to_cs in enumerate(bear_headers[1:4], 1):
            if to_cs and to_cs not in ('NOTHING HEARD', 'STATION SECURED') and col_idx < len(row) and row[col_idx] is not None:
                check(f"bearing({from_cs},{to_cs})", lookup_bearing(from_cs, to_cs, bear_matrix), float(row[col_idx]))

    # 4. Sector count
    print("\n4. Sector Count (COUNTIF):")
    mock = [{'sector':'CHARLESTON'},{'sector':'JACKSONVILLE'},{'sector':'CHARLESTON'},
            {'sector':'MIAMI'},{'sector':'JACKSONVILLE'},{'sector':'MIAMI'},{'sector':'MIAMI'}]
    counts = count_by_sector(mock)
    check("COUNTIF CHARLESTON",  counts['CHARLESTON'],  2)
    check("COUNTIF JACKSONVILLE",counts['JACKSONVILLE'], 2)
    check("COUNTIF MIAMI",       counts['MIAMI'],        3)
    check("SUM total",           counts['TOTAL'],        7)

    # 5. Frequency formula
    print("\n5. Frequency Display Formula:")
    check("format_freq A2E", format_freq_display("A2E"), "A2E:   2.8103 MHz")
    check("format_freq D7A", format_freq_display("D7A"), "D7A:   7.7365 MHz")

    # 6. DB vs Excel Email List cross-check
    print("\n6. DB vs Excel Email List Cross-Check:")
    active_cs = [cs for cs in excel_members if cs not in ('NOTHING HEARD','STATION SECURED')]
    db_missing = []
    field_mismatches = []
    conn = get_db()
    for cs in active_cs:
        db_row = conn.execute("SELECT * FROM members WHERE UPPER(callsign)=?", (cs.upper(),)).fetchone()
        if not db_row:
            db_missing.append(cs)
            tests_run += 1
            tests_failed += 1
        else:
            tests_run += 1
            tests_passed += 1
            excel_name = excel_members[cs]['name'] or ''
            db_name = db_row['operator_name'] or ''
            if excel_name.strip() != db_name.strip():
                field_mismatches.append(f"  name mismatch {cs}: DB={db_name!r} Excel={excel_name!r}")

    db_count = conn.execute("SELECT COUNT(*) FROM members WHERE is_synthetic=0").fetchone()[0]
    if db_missing:
        print(f"  [FAIL] DB missing callsigns: {db_missing}")
    else:
        print(f"  [PASS] All {len(active_cs)} Excel callsigns found in DB")
    if field_mismatches:
        for m in field_mismatches:
            print(f"  [WARN] {m}")
    print(f"  [INFO] DB has {db_count} active members | Excel has {len(active_cs)} active members")

    print(f"\n{'='*50}")
    print(f"Results: {tests_passed}/{tests_run} passed, {tests_failed} failed")
    if tests_failed == 0:
        print("All tests PASSED — Python DB output matches Excel exactly.")
    else:
        print(f"WARNING: {tests_failed} mismatches found.")


# ─────────────────────────────────────────────────────────────────────────────
# INTERACTIVE CLI
# ─────────────────────────────────────────────────────────────────────────────

def _prompt(label: str, default: str = "", choices: list = None, required: bool = False) -> str:
    while True:
        if choices:
            print(f"\n  {label}:")
            for i, c in enumerate(choices, 1):
                print(f"    {i:2}. {c}")
            raw = input(f"  Enter number or value [{default}]: ").strip()
            if not raw:
                return default
            if raw.isdigit() and 1 <= int(raw) <= len(choices):
                return choices[int(raw) - 1]
            if raw in choices:
                return raw
            print("  Invalid choice.")
        else:
            raw = input(f"  {label} [{default}]: ").strip() if default else input(f"  {label}: ").strip()
            if raw:
                return raw
            if default:
                return default
            if required:
                print("  Required — please enter a value.")
            else:
                return ""


def print_session_header(session: NetSession):
    print("\n" + "="*70)
    print(" USCG AUXILIARY SOUTHEAST HF CONTINGENCY NET LOG")
    print(f" Freq Type: {session.freq_type}  |  Net: {session.net_name}")
    print(f" Date: {session.net_date}  |  Start: {session.start_time}  |  NCS: {session.ncs_callsign}")
    print(f" Primary: {session.primary_freq}  |  Transceiver: {session.transceiver}")
    print(f" Location: {session.location}  |  Antenna: {session.antenna}")
    print(f" Show My D&B: {'YES' if session.show_my_db else 'NO'}  |  Manual Time: {'YES' if session.manual_time else 'NO'}")
    print("="*70)


def print_checkins_table(session: NetSession):
    counts = session.sector_counts
    print(f"\n{'#':>3} {'FROM':12} {'TO':12} {'TIME':10} {'SIG':15} {'CITY':18} {'SECTOR':16} {'DIST':8} {'BRG':7}")
    print("-"*105)
    for ci in session.checkins:
        sig = " ".join(c for c in ['L','G','W','VW','F','C','R','UR','D','WI','I'] if ci.get(f'signal_{c}'))
        dist = ci.get('distance_miles', '')
        brg  = ci.get('bearing_degrees', '')
        dist_str = f"{dist:.1f}" if isinstance(dist, float) else str(dist)[:8]
        brg_str  = f"{brg:.1f}°" if isinstance(brg, float) else str(brg)[:7]
        print(f"{ci.get('checkin_order', ''):>3} {ci.get('from_callsign',''):12} {ci.get('to_callsign',''):12} "
              f"{ci.get('checkin_time',''):10} {sig:15} {(ci.get('city','') or '')[:18]:18} "
              f"{(ci.get('sector','') or '')[:16]:16} {dist_str:8} {brg_str:7}")
    print("\nSector Counts: " + " | ".join(f"{s[:6]}:{counts[s]}" for s in SECTORS if counts[s] > 0))
    print(f"Total Checked-In: {counts['TOTAL']}")


def run_interactive(session: NetSession, dist_matrix: dict, bear_matrix: dict):
    print_session_header(session)
    if session.checkins:
        print_checkins_table(session)

    MENU = ("Commands: [A]dd check-in  [V]iew log  [E]dit session  [S]ave  [R]eset\n"
            "          [T]oggle D&B  [M]anual time  [B]rowse members  [N]ew member\n"
            "          [X] Edit member  [D]elete member  [Q]uit/Archive")

    while True:
        print("\n" + "-"*40)
        print(MENU)
        cmd = input("Command: ").strip().upper()

        if cmd == 'A':
            print("\n--- NEW CHECK-IN ---")
            from_cs = input("  FROM callsign: ").strip().upper()
            if not from_cs:
                continue
            op     = lookup_operator(from_cs)
            city   = lookup_city(from_cs)
            sector = lookup_sector(from_cs)
            print(f"  -> {op}, {city}, {sector}")

            to_cs = input("  TO callsign [ENTER for blank]: ").strip().upper()

            if session.manual_time:
                checkin_time = input("  Check-in time (HH:MM:SS): ").strip()
            else:
                checkin_time = datetime.now().strftime('%H:%M:%S')
                print(f"  Auto-timestamp: {checkin_time}")

            print("  Signal report (space-separated codes: L G W VW F C R UR D WI I)")
            sig_input = input("  Signal codes: ").strip().upper().split()
            signals = {f'signal_{c}': (c in sig_input) for c in ['L','G','W','VW','F','C','R','UR','D','WI','I']}

            has_traffic = input("  Has traffic? (y/N): ").strip().lower() == 'y'
            traffic_dest = traffic_notes = ""
            if has_traffic:
                traffic_dest  = input("  Traffic destination callsign: ").strip().upper()
                traffic_notes = input("  Traffic notes: ").strip()

            ci = session.add_checkin(
                dist_matrix, bear_matrix,
                from_callsign=from_cs, to_callsign=to_cs,
                checkin_time=checkin_time, manual_time=session.manual_time,
                has_traffic=has_traffic, traffic_destination=traffic_dest,
                traffic_notes=traffic_notes, **signals
            )
            dist = ci.distance_miles
            brg  = ci.bearing_degrees
            print(f"\n  Added: #{ci.checkin_order} {from_cs} -> {to_cs} @ {checkin_time}")
            print(f"  Distance: {f'{dist:.1f} mi' if isinstance(dist,(int,float)) else dist}"
                  f"  |  Bearing: {f'{brg:.1f}°' if isinstance(brg,(int,float)) else brg}")

        elif cmd == 'V':
            print_session_header(session)
            if session.checkins:
                print_checkins_table(session)
            else:
                print("No check-ins yet.")

        elif cmd == 'E':
            print("\n--- EDIT SESSION SETTINGS ---")
            session.ncs_callsign = input(f"  NCS callsign [{session.ncs_callsign}]: ").strip().upper() or session.ncs_callsign
            session.propagation  = _prompt("Propagation", session.propagation, PROPAGATIONS)
            session.noise_level  = _prompt("Noise Level", session.noise_level, NOISE_LEVELS)
            session.secured_time = input(f"  Secured time [{session.secured_time}]: ").strip() or session.secured_time
            session.recompute_all(dist_matrix, bear_matrix)
            print("  Session updated and formulas recomputed.")

        elif cmd == 'T':
            session.show_my_db = not session.show_my_db
            session.recompute_all(dist_matrix, bear_matrix)
            print(f"  Show My D&B: {'ON' if session.show_my_db else 'OFF'} — distances/bearings recomputed")

        elif cmd == 'M':
            session.manual_time = not session.manual_time
            print(f"  Manual time entry: {'ON' if session.manual_time else 'OFF'}")

        elif cmd == 'S':
            save_path = input("  Save to [netlog_session.json]: ").strip() or "netlog_session.json"
            session.save(save_path)
            print(f"  Saved to {save_path}")

        elif cmd == 'R':
            if input("  Reset all check-ins? (yes/NO): ").strip().lower() == 'yes':
                session.checkins = []
                print("  Check-ins cleared.")

        elif cmd == 'B':
            _cmd_browse_members()

        elif cmd == 'N':
            _cmd_new_member()

        elif cmd == 'X':
            _cmd_edit_member()

        elif cmd == 'D':
            _cmd_delete_member()

        elif cmd == 'Q':
            if input("  Archive this session? (y/N): ").strip().lower() == 'y':
                save_to_archive(session)
            break

        else:
            print("  Unknown command.")


def new_session_wizard(dist_matrix: dict, bear_matrix: dict) -> NetSession:
    print("\n+==================================================================+")
    print("|  USCG AUXILIARY D7 HF CONTINGENCY NET LOG  --  New Session      |")
    print("+==================================================================+\n")

    session = NetSession()
    session.freq_type   = _prompt("Frequency type", "PRIMARY", ["PRIMARY", "ALT", "2ndALT"])
    session.net_name    = _prompt("Net name", NET_NAMES[0], NET_NAMES)
    session.net_date    = input(f"  Net date (YYYY-MM-DD) [{date.today()}]: ").strip() or str(date.today())
    session.start_time  = input("  Start time (HHMM): ").strip()

    session.ncs_callsign = input("  NCS / My callsign: ").strip().upper()
    if session.ncs_callsign:
        op = lookup_operator(session.ncs_callsign)
        if op != 'Callsign FROM Not Found':
            print(f"  -> Recognized: {op}")
        else:
            print("  -> Not found in member DB (entry will still work)")

    session.transceiver = _prompt("Transceiver", "NONE", TRANSCEIVERS)
    session.antenna     = _prompt("Antenna", "ANTENNA #1", ANTENNAS)
    session.location    = _prompt("Location", "HOME", LOCATIONS)
    session.primary_freq = _prompt("Primary frequency", FREQ_DISPLAY[0], FREQ_DISPLAY)
    session.alt_freq_1  = _prompt("1st Alt frequency", FREQ_DISPLAY[0], FREQ_DISPLAY)
    session.alt_freq_2  = _prompt("2nd Alt frequency", FREQ_DISPLAY[0], FREQ_DISPLAY)
    session.digital_mode = _prompt("Digital mode", "NONE", OP_CODES)
    session.propagation = _prompt("Propagation", PROPAGATIONS[1], PROPAGATIONS)
    session.noise_level = _prompt("Noise level", NOISE_LEVELS[0], NOISE_LEVELS)

    session.show_my_db  = input("\n  Show My Station D&B to all stations? (y/N): ").strip().lower() == 'y'
    session.manual_time = input("  Manual check-in time entry? (y/N): ").strip().lower() == 'y'
    return session


def print_archives():
    archives = load_archives()
    if not archives:
        print("No archived sessions found.")
        return
    print(f"\n=== ARCHIVED SESSIONS ({len(archives)}) ===\n")
    for i, s in enumerate(archives, 1):
        print(f"{i:3}. {s.get('net_date','')} {s.get('net_name','')} | "
              f"Freq: {s.get('freq_type','')} | NCS: {s.get('ncs_callsign','')} | "
              f"Check-ins: {len(s.get('checkins',[]))} | Archived: {s.get('archived_at','')[:10]}")


# ─────────────────────────────────────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────────────────────────────────────

def main():
    args = sys.argv[1:]

    if '--archives' in args:
        print_archives()
        return

    xlsx_candidates = [
        os.path.join(os.path.dirname(os.path.abspath(__file__)), 'NetLog_decrypted.xlsx'),
        'NetLog_decrypted.xlsx',
    ]

    if '--verify' in args:
        idx = args.index('--verify')
        verify_path = args[idx+1] if idx+1 < len(args) else next((p for p in xlsx_candidates if os.path.exists(p)), None)
        if not verify_path:
            print("Provide path: python netlog.py --verify /path/to/NetLog_decrypted.xlsx")
            return
        verify_against_excel(verify_path)
        return

    xlsx_path = next((p for p in xlsx_candidates if os.path.exists(p)), None)
    dist_matrix, bear_matrix = {}, {}
    if xlsx_path:
        print(f"Loading D/B matrices from {xlsx_path}...")
        dist_matrix, bear_matrix = load_matrices(xlsx_path)
        print(f"  Loaded {len(dist_matrix)} distance pairs, {len(bear_matrix)} bearing pairs.")
    else:
        print("Note: NetLog_decrypted.xlsx not found — distance/bearing lookups unavailable.")
        print("  To enable: python netlog.py --verify /path/to/decrypted.xlsx")

    if '--load' in args:
        idx = args.index('--load')
        load_path = args[idx+1] if idx+1 < len(args) else None
        if not load_path or not os.path.exists(load_path):
            print(f"Session file not found: {load_path}")
            return
        session = NetSession.load(load_path)
        print(f"Loaded session from {load_path}")
    else:
        session = new_session_wizard(dist_matrix, bear_matrix)

    run_interactive(session, dist_matrix, bear_matrix)


if __name__ == '__main__':
    main()
